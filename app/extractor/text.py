"""Extract selectable text from uploaded PDF and DOCX files."""

from io import BytesIO
from pathlib import Path
import pdfplumber
from docx import Document

from app.utils import clean_text


class TextExtractionError(ValueError):
    pass


def extract_text(filename: str, content: bytes) -> str:
    suffix = Path(filename or "").suffix.lower()
    if suffix not in {".pdf", ".docx"}:
        raise TextExtractionError("Unsupported file type. Upload a PDF or DOCX file.")
    if not content:
        raise TextExtractionError("The uploaded file is empty.")

    try:
        if suffix == ".pdf":
            with pdfplumber.open(BytesIO(content)) as pdf:
                parts = [page.extract_text(layout=False) or "" for page in pdf.pages]
                education_rows = []
                for page in pdf.pages:
                    # Ruled PDF tables retain their cell boundaries here even when
                    # extract_text() flattens all columns into one long line.
                    for table in page.extract_tables():
                        if not table or not any(_is_education_header(row) for row in table[:2]):
                            continue
                        education_rows.extend(
                            " | ".join((cell or "").replace("\n", " ").strip() for cell in row)
                            for row in table
                            if row and not _is_education_header(row)
                        )
                if education_rows:
                    parts.append("Education\n" + "\n".join(education_rows))
        else:
            doc = Document(BytesIO(content))
            parts = [paragraph.text for paragraph in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
    except Exception as exc:
        # Parsing libraries raise several format-specific errors for damaged files.
        raise TextExtractionError(f"Could not read the {suffix[1:].upper()} file: {exc}") from exc

    text = clean_text("\n".join(parts))
    if not text:
        raise TextExtractionError("No selectable text was found. Scanned PDFs need OCR, which this app does not provide.")
    return text


def _is_education_header(row: list[str | None]) -> bool:
    cells = [(cell or "").strip().casefold() for cell in row]
    return "year" in cells and "degree" in cells and "institute" in cells


def extract_jd_file(filename: str, content: bytes) -> str:
    """A JD can be plain UTF-8 text as well as PDF or DOCX."""
    if Path(filename or "").suffix.lower() != ".txt":
        return extract_text(filename, content)
    if not content:
        raise TextExtractionError("The uploaded job description is empty.")
    try:
        text = clean_text(content.decode("utf-8-sig"))
    except UnicodeDecodeError as exc:
        raise TextExtractionError("The job description text file must be UTF-8 encoded.") from exc
    if not text:
        raise TextExtractionError("The job description text file is empty.")
    return text
